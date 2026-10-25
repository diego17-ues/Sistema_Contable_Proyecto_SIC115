import argparse
import os
import re
import sys
from pathlib import Path

import docx
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader

load_dotenv()
MANUALES = [Path(p.strip()) for p in os.getenv("CARPETA_MANUALES", "manuales").split(";") if p.strip()]
REQUERIMIENTOS = Path(os.getenv("CARPETA_REQUERIMIENTOS", "requerimientos"))
SALIDA = Path(os.getenv("CARPETA_SALIDA", str(REQUERIMIENTOS / "generados")))
MODELO = os.environ["OPENAI_MODEL"]
EXTENSIONES = {".md", ".txt", ".docx", ".pdf"}
client = OpenAI()


def leer(ruta: Path) -> str:
    ext = ruta.suffix.lower()
    if ext in (".md", ".txt"):
        return ruta.read_text(encoding="utf-8")
    if ext == ".docx":
        d = docx.Document(str(ruta))
        partes = [p.text for p in d.paragraphs if p.text.strip()]
        for tabla in d.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)
    if ext == ".pdf":
        return "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)
    return ""


def cargar_carpeta(carpeta: Path, etiqueta: str, excluir: Path | None = None) -> str:
    if not carpeta.is_dir():
        sys.exit(f"No existe la carpeta: {carpeta}")
    bloques = []
    for ruta in sorted(carpeta.rglob("*")):
        if not ruta.is_file() or ruta.suffix.lower() not in EXTENSIONES:
            continue
        if ruta.name.startswith("~$"):          # temporales de Word
            continue
        if excluir and excluir.resolve() in ruta.resolve().parents:
            continue                             # no releer lo ya generado
        texto = leer(ruta)
        if not texto.strip():
            print(f"AVISO: no pude leer texto de {ruta}")
            continue
        nombre = f"{carpeta.name}/{ruta.relative_to(carpeta).as_posix()}"  
        bloques.append(f'<{etiqueta} nombre="{nombre}">\n{texto}\n</{etiqueta}>')
    return "\n\n".join(bloques)


def llamar(instrucciones: str, entrada: str) -> str:
    resp = client.responses.create(
        model=MODELO, instructions=instrucciones, input=entrada
    )
    u = resp.usage
    print(f"  tokens: entrada={u.input_tokens}, salida={u.output_tokens}")
    return resp.output_text


TAREAS = {
    "redactar": (
        "Redacta los requerimientos funcionales y no funcionales del módulo: "
        "{modulo}. Basa los hechos en los manuales. Usa los requerimientos "
        "existentes solo como referencia de estilo y numeración, sin duplicarlos "
        "ni contradecirlos."
    ),
    "ajustar": (
        "Reajusta los requerimientos existentes correspondientes a: {modulo}, "
        "para que respeten el alcance y los manuales. Entrega primero la versión "
        "ajustada y al final una 'Bitácora de cambios' que indique, por cada ID "
        "original, si fue CONSERVADO, MODIFICADO, ELIMINADO (con motivo) o si se "
        "agregó uno NUEVO."
    ),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("modulo", help='Ej: "contabilidad general" o "todo el documento"')
    ap.add_argument("--modo", choices=TAREAS, default="redactar")
    ap.add_argument("--revisar", action="store_true", help="Pasa un revisor al borrador")
    args = ap.parse_args()

    manuales = "\n\n".join(cargar_carpeta(c, "manual") for c in MANUALES)
    existentes = cargar_carpeta(REQUERIMIENTOS, "requerimiento_existente", excluir=SALIDA)
    if not manuales:
        sys.exit(f"No encontré texto legible en {MANUALES}")
    if args.modo == "ajustar" and not existentes:
        sys.exit(f"El modo 'ajustar' necesita archivos en {REQUERIMIENTOS}")

    entrada = (
        f"<manuales>\n{manuales}\n</manuales>\n\n"
        f"<requerimientos_existentes>\n{existentes or '(ninguno)'}\n</requerimientos_existentes>\n\n"
        f"Tarea: {TAREAS[args.modo].format(modulo=args.modulo)}"
    )

    nombre = re.sub(r"\W+", "_", args.modulo.lower()).strip("_")
    SALIDA.mkdir(parents=True, exist_ok=True)

    print("Redactando...")
    borrador = llamar(Path("instrucciones.md").read_text(encoding="utf-8"), entrada)
    destino = SALIDA / f"req_{nombre}_{args.modo}.md"
    destino.write_text(borrador, encoding="utf-8")
    print(f"Borrador: {destino}")

    if args.revisar:
        print("Revisando...")
        contexto = f"<manuales>\n{manuales}\n</manuales>\n\n"
        if args.modo == "ajustar":
            contexto += f"<requerimientos_existentes>\n{existentes}\n</requerimientos_existentes>\n\n"
        informe = llamar(
            Path("instrucciones_revisor.md").read_text(encoding="utf-8"),
            f"{contexto}<borrador>\n{borrador}\n</borrador>",
        )
        destino_rev = SALIDA / f"req_{nombre}_{args.modo}_revision.md"
        destino_rev.write_text(informe, encoding="utf-8")
        print(f"Revisión: {destino_rev}")


main()