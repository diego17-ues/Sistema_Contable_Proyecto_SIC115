import argparse
import json
import os
import sys
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader

load_dotenv()
MANUALES = [Path(p.strip()) for p in os.getenv("CARPETA_MANUALES", "manuales").split(";") if p.strip()]
SALIDA = Path(os.getenv("CARPETA_SALIDA", "salida"))
LOTES = SALIDA / "lotes_ieee"
MODELO = os.environ["OPENAI_MODEL"]
PRECIO_IN = float(os.getenv("PRECIO_ENTRADA", "0"))   # USD por millón de tokens
PRECIO_OUT = float(os.getenv("PRECIO_SALIDA", "0"))
TOPE_USD = float(os.getenv("TOPE_USD", "0"))           # 0 = sin tope
EXTENSIONES = {".md", ".txt", ".docx", ".pdf"}
CLAVE = ["Manual_aplicación_catalogo_SunGames.docx", "Modelo_de_Negocio_SunGames.docx"]
TOTAL = {"in": 0, "out": 0}
client = OpenAI()

# ---- Qué se genera: (nombre, cantidad, alcance). Edita libremente. ----
MODULOS_RF = [
    ("Usuarios, roles y acceso", 4, "alta, baja y bloqueo de usuarios, asignación de rol (administrador y contador; costeador solo si los manuales lo respaldan), inicio y cierre de sesión, cambio de contraseña"),
    ("Catálogo de cuentas", 6, "alta, modificación, desactivación y consulta de cuentas según el manual de aplicación del catálogo, rubros, código y naturaleza, sin subcuentas, carga inicial"),
    ("Parámetros y períodos contables", 5, "datos de la empresa, ejercicio y períodos contables, apertura y cierre de período, moneda, correlativos y parámetros fiscales que respalden los manuales"),
    ("Partidas y libro diario", 7, "registro, validación de partida doble, edición antes de mayorizar, anulación y reversión, partidas modelo, consulta y filtros del libro diario"),
    ("Libro mayor y balance de comprobación", 5, "mayorización, saldos por cuenta, consulta por rango de fechas, balance de comprobación, exportación"),
    ("Ingresos por venta de licencias", 6, "registro de ventas digitales de pago único por Steam, comisión de la plataforma, impuestos según los manuales, conciliación de liquidaciones, partida automática"),
    ("Egresos, compras y gastos", 5, "registro de compras y gastos de desarrollo, proveedores, cuentas por pagar, pagos, partida automática"),
    ("Inventario", 5, "lo que los manuales definan como inventario o activos de la empresa, entradas, salidas, valuación y consulta de existencias"),
    ("Contabilidad de costos por órdenes de trabajo", 8, "apertura de orden, asignación de materiales, mano de obra y costos indirectos, aplicación de costos indirectos, cierre de orden, costo unitario y total del videojuego, integración con la contabilidad general"),
    ("Reportes financieros", 6, "balance general, estado de resultados, estado de cambios en el patrimonio, reportes de costos, filtros por período, exportación a PDF y Excel"),
    ("Ajustes y cierre contable", 4, "partidas de ajuste, cierre de período y de ejercicio, reapertura controlada"),
    ("Bitácora y respaldo", 3, "consulta y exportación de la bitácora de operaciones, respaldo y restauración de datos"),
]
MODULOS_RNF = [
    ("Seguridad y privacidad", 14, "autenticación, contraseñas, bloqueo por intentos, sesiones, autorización por rol, cifrado en tránsito y en reposo, protección contra inyección y XSS, protección de datos"),
    ("Integridad, auditoría y trazabilidad", 14, "partida doble siempre cuadrada, inmutabilidad de partidas contabilizadas, correlativos sin saltos, bitácora de cambios, conservación de registros, consistencia transaccional"),
    ("Rendimiento, escalabilidad y concurrencia", 12, "tiempos de respuesta, tiempos de reportes, volumen de partidas, usuarios simultáneos, bloqueos y concurrencia, crecimiento de datos"),
    ("Disponibilidad, respaldo y recuperación", 12, "disponibilidad, ventanas de mantenimiento, respaldos, restauración, tolerancia a fallos, manejo de errores"),
    ("Usabilidad, accesibilidad y localización", 12, "facilidad de uso para un contador, mensajes de error claros, navegación, accesibilidad, diseño adaptable, español, moneda USD, formato de fechas y números"),
    ("Mantenibilidad, portabilidad e interoperabilidad", 12, "arquitectura y código, pruebas, documentación, versionado, despliegue, compatibilidad de navegadores y bases de datos, exportación e importación de datos"),
    ("Cumplimiento legal y normativo", 8, "conservación de libros y registros, normativa mercantil, tributaria y contable de El Salvador, solo con respaldo de los manuales"),
]
TAREA = {
    "rf": ("Tarea: genera EXACTAMENTE {n} requerimientos funcionales para el submódulo «{nombre}». "
           "Alcance a cubrir: {alcance}. Un requerimiento por cada función, validación, regla, error, "
           "permiso o consulta. Ya existen estos en otros submódulos (no los repitas): {previos}."),
    "rnf": ("Tarea: genera EXACTAMENTE {n} requerimientos NO funcionales del grupo «{nombre}». "
            "Temas: {alcance}. Pueden ser generales de cualquier sistema contable, cada uno con una "
            "métrica o criterio verificable. Ya existen estos (no los repitas): {previos}."),
}

# ---- Esquema del JSON ----
LISTA = {"type": "array", "items": {"type": "string"}}
PROPS = {
    "nombre": {"type": "string"}, "descripcion": {"type": "string"},
    "prioridad": {"type": "string", "enum": ["Alta", "Media", "Baja"]},
    "actores": LISTA, "fuente": {"type": "string"}, "entradas": LISTA,
    "validacion_entradas": LISTA, "validacion_salidas": LISTA, "secuencia": LISTA,
    "situaciones_anormales": LISTA, "parametros": LISTA, "salidas": LISTA,
    "almacenamiento_bd": LISTA,
}
REQ = {"type": "object", "additionalProperties": False, "required": list(PROPS), "properties": PROPS}
SCHEMA = {"type": "object", "additionalProperties": False, "required": ["requerimientos", "vacios"],
          "properties": {"requerimientos": {"type": "array", "items": REQ}, "vacios": LISTA}}
FILAS = [
    ("Identificador", "identificador"), ("Nombre del requerimiento", "nombre"),
    ("Descripción", "descripcion"), ("Prioridad", "prioridad"), ("Actores", "actores"),
    ("Fuente del requisito", "fuente"), ("Entradas requeridas", "entradas"),
    ("Validación de entradas", "validacion_entradas"), ("Validación de salidas", "validacion_salidas"),
    ("Secuencia de operaciones", "secuencia"),
    ("Respuesta a situaciones anormales", "situaciones_anormales"),
    ("Parámetros del sistema", "parametros"), ("Salidas generadas", "salidas"),
    ("Requerimientos lógicos de almacenamiento en BD", "almacenamiento_bd"),
]


def gastado():
    return TOTAL["in"] / 1e6 * PRECIO_IN + TOTAL["out"] / 1e6 * PRECIO_OUT


# ---- Lectura de documentos ----
def leer(ruta: Path) -> str:
    ext = ruta.suffix.lower()
    if ext in (".md", ".txt"):
        return ruta.read_text(encoding="utf-8")
    if ext == ".docx":
        d = Document(str(ruta))
        partes = [p.text for p in d.paragraphs if p.text.strip()]
        for tabla in d.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)
    if ext == ".pdf":
        return "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)
    return ""


def norm(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


def buscar(nombre: str):
    for carpeta in MANUALES:
        if carpeta.is_dir():
            for ruta in carpeta.rglob("*"):
                if ruta.is_file() and norm(ruta.name) == norm(nombre):
                    return ruta
    return None


def bloque(ruta: Path, principal: bool = False):
    texto = leer(ruta)
    if not texto.strip():
        print(f"AVISO: no pude leer texto de {ruta.name}")
        return None
    print(f"  {ruta.name}: ~{len(texto) // 4} tokens")
    marca = ' principal="si"' if principal else ""
    return f'<manual nombre="{ruta.name}"{marca}>\n{texto}\n</manual>'


def cargar_contexto(todo: bool) -> str:
    rutas, bloques = [], []
    for nombre in CLAVE:
        ruta = buscar(nombre)
        if not ruta:
            sys.exit(f"No encontré {nombre} en {[str(c) for c in MANUALES]}")
        b = bloque(ruta, True)
        if b is None:
            sys.exit(f"No se pudo leer {nombre}")
        rutas.append(ruta.resolve())
        bloques.append(b)
    if todo:
        for carpeta in MANUALES:
            for ruta in sorted(carpeta.rglob("*")):
                if (ruta.is_file() and ruta.suffix.lower() in EXTENSIONES
                        and not ruta.name.startswith("~$") and ruta.resolve() not in rutas):
                    b = bloque(ruta)
                    if b:
                        bloques.append(b)
    return "\n\n".join(bloques)


# ---- Llamadas a la API ----
def pedir(instrucciones: str, contexto: str, tarea: str) -> dict:
    for _ in range(2):
        resp = client.responses.create(
            model=MODELO,
            instructions=instrucciones,
            input=f"{contexto}\n\n{tarea}",
            text={"format": {"type": "json_schema", "name": "lote", "schema": SCHEMA, "strict": True}},
        )
        u = resp.usage
        TOTAL["in"] += u.input_tokens
        TOTAL["out"] += u.output_tokens
        cache = getattr(getattr(u, "input_tokens_details", None), "cached_tokens", 0) or 0
        print(f"  tokens: entrada={u.input_tokens} (caché={cache}), salida={u.output_tokens}")
        if resp.status == "incomplete":
            sys.exit("Respuesta cortada por límite de salida. Baja la cantidad de ese lote en el script.")
        try:
            return json.loads(resp.output_text)
        except json.JSONDecodeError:
            print("  JSON inválido, reintentando...")
    sys.exit("El modelo no devolvió JSON válido dos veces.")


def generar(tipo, modulos, instrucciones, contexto, rehacer, solo_lote=None):
    LOTES.mkdir(parents=True, exist_ok=True)
    previos = []
    for i, (nombre, n, alcance) in enumerate(modulos, 1):
        archivo = LOTES / f"{tipo}_{i:02d}.json"
        if solo_lote and i != solo_lote:
            if archivo.exists():
                previos.extend(r["nombre"] for r in json.loads(archivo.read_text(encoding="utf-8"))["requerimientos"])
            continue
        if archivo.exists() and not rehacer:
            datos = json.loads(archivo.read_text(encoding="utf-8"))
            print(f"[{tipo.upper()} {i}/{len(modulos)}] {nombre}: reutilizado ({len(datos['requerimientos'])})")
        else:
            print(f"[{tipo.upper()} {i}/{len(modulos)}] {nombre}...")
            tarea = TAREA[tipo].format(n=n, nombre=nombre, alcance=alcance, previos="; ".join(previos) or "(ninguno)")
            datos = pedir(instrucciones, contexto, tarea)
            archivo.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
            print(f"  generados: {len(datos['requerimientos'])}")
            if PRECIO_IN:
                print(f"  gasto estimado acumulado: ${gastado():.3f}")
            if TOPE_USD and gastado() >= TOPE_USD:
                sys.exit(f"Tope de ${TOPE_USD} alcanzado. Lo generado ya está guardado; sube el tope y vuelve a correr para continuar.")
        previos.extend(r["nombre"] for r in datos["requerimientos"])


def cargar_lotes(tipo, modulos):
    salida = []
    for i, (nombre, _, _) in enumerate(modulos, 1):
        archivo = LOTES / f"{tipo}_{i:02d}.json"
        if archivo.exists():
            salida.append((nombre, json.loads(archivo.read_text(encoding="utf-8"))))
    return salida


# ---- Documento Word ----
def sombrear(celda, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(shd)


def escribir(celda, valor, numerar=False):
    if isinstance(valor, list):
        lineas = [f"{i}. {t}" if numerar else f"• {t}" for i, t in enumerate(valor, 1)] or ["No aplica"]
    else:
        lineas = [valor or "No aplica"]
    celda.text = ""
    for k, linea in enumerate(lineas):
        p = celda.paragraphs[0] if k == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(linea).font.size = Pt(9.5)


def tabla_requerimiento(doc, codigo, req):
    doc.add_heading(f"{codigo} | {req['nombre']}", level=3).paragraph_format.keep_with_next = True
    datos = {"identificador": codigo, **req}
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for etiqueta, clave in FILAS:
        fila = t.add_row()
        fila.cells[0].width, fila.cells[1].width = Cm(4.6), Cm(12.0)
        fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        c0 = fila.cells[0]
        c0.text = ""
        r = c0.paragraphs[0].add_run(etiqueta)
        r.bold, r.font.size = True, Pt(9.5)
        sombrear(c0, "D9E2F3")
        escribir(fila.cells[1], datos[clave], numerar=(clave == "secuencia"))
    doc.add_paragraph()


def indice(doc):
    r = doc.add_paragraph().add_run()
    partes = []
    for tipo in ("begin", None, "separate", "txt", "end"):
        if tipo is None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = 'TOC \\o "1-3" \\h \\z \\u'
        elif tipo == "txt":
            el = OxmlElement("w:t")
            el.text = "Clic derecho aquí y elige «Actualizar campo» para generar el índice."
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        partes.append(el)
    for el in partes:
        r._r.append(el)


def construir(rf, rnf):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_before = Pt(140)
    for texto, tam, negrita in (("Especificación de requerimientos de software", 24, True),
                                ("Sistema de contabilidad general y de costos", 16, False),
                                ("SunGames, S.A. de C.V.", 16, False),
                                (f"Borrador generado el {date.today():%d/%m/%Y}", 11, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold, r.font.size = negrita, Pt(tam)
    doc.add_page_break()
    doc.add_heading("Índice", level=1)
    indice(doc)
    doc.add_page_break()

    doc.add_heading("1. Introducción", level=1)
    doc.add_paragraph(
        "Este documento especifica los requerimientos funcionales y no funcionales del sistema "
        "contable de SunGames, S.A. de C.V. Cada requerimiento se presenta en una tabla con los "
        "elementos propios de una especificación de requisitos (entradas, validaciones, secuencia "
        "de operaciones, situaciones anormales, parámetros, salidas y almacenamiento).")
    doc.add_paragraph(
        "Fuentes principales: Manual de aplicación del catálogo de cuentas y Modelo de negocio de "
        "SunGames. Borrador elaborado con apoyo de IA, sujeto a revisión del equipo.")

    vacios, n_rf, n_rnf = [], 0, 0
    doc.add_heading("2. Requerimientos funcionales", level=1)
    for i, (sub, datos) in enumerate(rf, 1):
        doc.add_heading(f"2.{i} {sub}", level=2)
        for req in datos["requerimientos"]:
            n_rf += 1
            tabla_requerimiento(doc, f"RF-{n_rf:03d}", req)
        vacios += datos["vacios"]
    doc.add_heading("3. Requerimientos no funcionales", level=1)
    for i, (grupo, datos) in enumerate(rnf, 1):
        doc.add_heading(f"3.{i} {grupo}", level=2)
        for req in datos["requerimientos"]:
            n_rnf += 1
            tabla_requerimiento(doc, f"RNF-{n_rnf:03d}", req)
        vacios += datos["vacios"]
    doc.add_heading("4. Vacíos y preguntas abiertas", level=1)
    for v in dict.fromkeys(vacios):
        doc.add_paragraph(v, style="List Bullet")
    return doc, n_rf, n_rnf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--solo", choices=["rf", "rnf"], help="genera solo esa parte")
    ap.add_argument("--lote", type=int, help="genera solo ese número de lote (úsalo junto con --solo)")
    ap.add_argument("--rehacer", action="store_true", help="vuelve a generar los lotes ya guardados")
    ap.add_argument("--todo", action="store_true", help="incluye también el resto de documentos")
    ap.add_argument("--sin-generar", action="store_true", help="solo arma el Word con los lotes guardados")
    args = ap.parse_args()

    if not args.sin_generar:
        instrucciones = Path("instrucciones_ieee.md").read_text(encoding="utf-8")
        print("Documentos leídos:")
        contexto = cargar_contexto(args.todo)
        if args.solo != "rnf":
            generar("rf", MODULOS_RF, instrucciones, contexto, args.rehacer, args.lote)
        if args.solo != "rf":
            generar("rnf", MODULOS_RNF, instrucciones, contexto, args.rehacer, args.lote)
        print(f"Tokens totales de esta corrida: entrada={TOTAL['in']}, salida={TOTAL['out']}")
        if PRECIO_IN:
            print(f"Gasto estimado de esta corrida: ${gastado():.3f}")

    doc, n_rf, n_rnf = construir(cargar_lotes("rf", MODULOS_RF), cargar_lotes("rnf", MODULOS_RNF))
    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / "Especificacion_Requerimientos_SunGames.docx"
    doc.save(destino)
    print(f"Word: {destino}  |  RF: {n_rf}  |  RNF: {n_rnf}")
    if n_rf <= 50:
        print("AVISO: faltan RF para superar 50. Sube las cantidades en MODULOS_RF.")
    if n_rnf <= 70:
        print("AVISO: faltan RNF para superar 70. Sube las cantidades en MODULOS_RNF.")


main()